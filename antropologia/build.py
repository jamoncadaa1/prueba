from PIL import Image, ImageDraw, ImageFilter
import random, math
random.seed(7)
W,H=1600,1000
img=Image.new("RGB",(W,H),(20,14,10))
d=ImageDraw.Draw(img)
for _ in range(9000):
    x,y=random.randint(0,W),random.randint(0,H)
    c=random.randint(28,70); d.point((x,y),fill=(c,int(c*.8),int(c*.6)))
glow=Image.new("L",(W,H),0); g=ImageDraw.Draw(glow)
for r in range(700,0,-10):
    g.ellipse((420-r,820-r,420+r,820+r),fill=int(200*(1-r/700)))
glow=glow.filter(ImageFilter.GaussianBlur(40))
warm=Image.new("RGB",(W,H),(255,150,60))
img=Image.composite(warm,img,glow.point(lambda v:int(v*.55)))
d=ImageDraw.Draw(img)
red=(165,40,28)
cx,cy=1000,420
d.ellipse((cx-90,cy-45,cx+90,cy+45),fill=red)
d.polygon([(cx-60,cy),(cx-420,cy-260),(cx-470,cy-150),(cx-380,cy-60),(cx-440,cy+20),(cx-120,cy+30)],fill=red)
d.polygon([(cx+60,cy),(cx+420,cy-260),(cx+470,cy-150),(cx+380,cy-60),(cx+440,cy+20),(cx+120,cy+30)],fill=red)
d.polygon([(cx-40,cy+35),(cx+40,cy+35),(cx+70,cy+190),(cx,cy+150),(cx-70,cy+190)],fill=red)
d.ellipse((cx-55,cy-120,cx+55,cy-20),fill=red)
d.polygon([(cx-18,cy-40),(cx+18,cy-40),(cx,cy+5)],fill=(120,28,20))
for ex in (cx-24,cx+24):
    d.ellipse((ex-14,cy-92,ex+14,cy-64),fill=(245,200,70))
    for a in range(0,360,30):
        t=math.radians(a); d.line((ex+16*math.cos(t),cy-78+16*math.sin(t),ex+26*math.cos(t),cy-78+26*math.sin(t)),fill=(245,200,70),width=3)
d.ellipse((330,860,510,920),fill=(70,45,30))
for i in range(40):
    x=420+random.randint(-40,40); h=random.randint(60,170)
    d.polygon([(x-18,870),(x+18,870),(x+random.randint(-10,10),870-h)],fill=(255,random.randint(120,210),40))
img=img.filter(ImageFilter.GaussianBlur(1.2))
img.save("guacharo.png")

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
doc=Document()
s=doc.sections[0]
s.page_height=Cm(27.94); s.page_width=Cm(21.59)
s.top_margin=Cm(3); s.bottom_margin=Cm(3); s.left_margin=Cm(4); s.right_margin=Cm(2)
st=doc.styles["Normal"]; st.font.name="Arial"; st.font.size=Pt(12)
st.element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia',"Arial")
pf=st.paragraph_format; pf.line_spacing=1.5; pf.space_after=Pt(0); pf.first_line_indent=Cm(0)
def P(t="",bold=False,center=False,just=True,italic=False):
    p=doc.add_paragraph(); r=p.add_run(t); r.bold=bold; r.italic=italic
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else (WD_ALIGN_PARAGRAPH.JUSTIFY if just else WD_ALIGN_PARAGRAPH.LEFT)
    return p
def H(t): P(t,bold=True,center=True); P()
def br(): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# portada
P("ACTIVIDAD 1. HISTORIAS",bold=True,center=True)
P("EL FUEGO EN LAS SOMBRAS",bold=True,center=True)
for _ in range(5): P()
P("JOSÉ LEONARDO BORRERO",center=True)
P("JULIÁN MONCADA",center=True)
for _ in range(6): P()
P("Antropología",center=True)
for _ in range(6): P()
P("24 de septiembre de 2026",center=True)
br()

H("1. ANÁLISIS DE LA HISTORIA ORIGINAL")
P("a. ¿Qué elementos culturales están presentes?",bold=True)
for t in [
"El elemento central es la religión. El autor explica que las religiones dan sentido a las vivencias de una sociedad, fijan prohibiciones y reglas, y según la religión de cada grupo las personas interpretan por qué ocurren las cosas.",
"Aparecen los lugares sagrados, sean mezquitas, iglesias u otros espacios, a los que se acude para adorar y rezar a entes sobrenaturales que cambian de una religión a otra.",
"También están las prácticas funerarias, que varían con la religión y con la cultura en la que cada persona se cría. Un ejemplo son los entierros de los neandertales.",
"Por último se mencionan las túnicas y vestimentas, los adornos corporales, las imágenes, las figuras y los grabados en las paredes, además de los rituales y los mitos. Todos forman parte de las tradiciones culturales."]:
    P(t)
P()
P("b. ¿Qué valores, creencias o prácticas se reflejan?",bold=True)
for k,v in [
("Respeto por los muertos:","los entierros muestran una atención especial hacia las personas fallecidas."),
("Creencia en un posible más allá:","el autor la plantea, aunque no precisa en qué creencias concretas se apoya."),
("Valor de lo sagrado:","ciertos lugares, como los espacios subterráneos, adquirieron un significado especial."),
("Simbolismo:","grabados, imágenes y adornos se usaban para la adoración."),
("Tradición cultural:","las fórmulas artísticas y los códigos se mantuvieron durante mucho tiempo."),
("Organización social:","el texto relaciona rituales, mitos, valores y diferencias de estatus entre individuos.")]:
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("• "+k+" ").bold=True; p.add_run(v)
br()

H("2. NUEVA HISTORIA: EL FUEGO EN LAS SOMBRAS")
story=open("historia.txt",encoding="utf-8").read().strip().split("\n\n")
for t in story: P(t)
P()
doc.add_picture("guacharo.png",width=Cm(14))
doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
P("Figura 1. Ilustración de la primera pintura del guácharo a la luz de la lámpara de grasa. Imagen elaborada digitalmente para esta actividad.",center=True,italic=True)
br()

H("3. SOCIALIZACIÓN DE LA PROPUESTA")
for t in open("socializacion.txt",encoding="utf-8").read().strip().split("\n\n"): P(t)
doc.save("Actividad1_Historias_Borrero_Moncada.docx")
